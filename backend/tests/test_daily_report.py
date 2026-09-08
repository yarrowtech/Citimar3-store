"""Per-store Daily Operations export (src/daily_report.py) -- built from the
mongomock daily logs, rendered through the shared Phase C xlsx/pdf writers."""
from __future__ import annotations

import io
from datetime import date, time

import openpyxl

from db.models import TARGETS, next_id
from src import daily_dashboard_store, daily_report
from src.reports.dispatch import RENDERERS
from src.reports.models import ChartBlock, GaugeBlock, KpiGridBlock, TableBlock, TitleBlock


def _seed(db):
    db[TARGETS].insert_one({
        "_id": next_id(db, TARGETS), "store_code": "NM", "entry_date": "2026-08-20",
        "sales_target": 10_000.0, "net_sales": None, "remaining": None, "footfall": None,
        "nob": None, "atv": None, "rpv": None, "basket_size": None,
        "conversion_pct": None, "achievement_pct": None, "reason": "Heavy rain.",
    })
    for d, t, amt, qty in [
        (date(2026, 8, 20), time(11, 30), 500.0, 2.0),
        (date(2026, 8, 20), time(15, 0), 800.0, 3.0),
        (date(2026, 8, 21), time(12, 0), 1_200.0, 4.0),
    ]:
        daily_dashboard_store.add_bill_entry(db, "NM", d, t, amt, qty)
    daily_dashboard_store.add_footfall_entry(db, "NM", date(2026, 8, 20), time(11, 30), 40.0)
    daily_dashboard_store.add_nob_entry(db, "NM", date(2026, 8, 20), time(11, 30), 4.0)


def test_report_dates_are_sorted_union(db_session):
    _seed(db_session)
    assert daily_report.report_dates(db_session, "NM") == ["2026-08-20", "2026-08-21"]


def test_xlsx_payload_is_tables_only(db_session):
    _seed(db_session)
    payload = daily_report.build_daily_report_payload(db_session, "NM", include_visuals=False)
    kinds = [type(b) for b in payload.blocks]
    assert kinds == [TitleBlock, TableBlock, TableBlock]

    content = RENDERERS["xlsx"].render(payload)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Summary", "Date & Time Slot Wise", "Date Wise Store Summary"]
    assert not any(ws._images for ws in wb.worksheets)  # no charts anywhere

    datewise = wb["Date Wise Store Summary"]
    assert [c.value for c in datewise[1]] == daily_report._DATEWISE_COLUMNS
    assert datewise.max_row == 1 + 2  # header + one row per recorded date
    remarks_col = daily_report._DATEWISE_COLUMNS.index("Remarks") + 1
    assert datewise.cell(row=2, column=remarks_col).value == "Heavy rain."


def test_timeslot_table_repeats_whole_day_target(db_session):
    _seed(db_session)
    rows = daily_report.build_timeslot_rows(db_session, "NM", ["2026-08-20", "2026-08-21"])
    aug20 = [r for r in rows if r["Date"] == "2026-08-20"]
    assert len(aug20) == 2  # the 11:30 and 15:00 bills fall in two different slots
    # Achievement % on a slot row = that slot's net sales / the whole-day target
    slot_1130 = next(r for r in aug20 if r["Time Slot"].startswith("11"))
    assert slot_1130["Net Sales"] == 500.0
    assert slot_1130["Achievement %"] == 5.0  # 500 / 10_000
    assert slot_1130["Remaining"] == 9_500.0


def test_pdf_payload_has_visuals_and_renders(db_session, monkeypatch):
    monkeypatch.setattr(daily_report.daily_context, "get_weather", lambda _d: None)  # no network in tests
    _seed(db_session)
    payload = daily_report.build_daily_report_payload(db_session, "NM", include_visuals=True)
    kinds = [type(b) for b in payload.blocks]
    assert kinds.count(GaugeBlock) == 6
    assert ChartBlock in kinds and KpiGridBlock in kinds

    content = RENDERERS["pdf"].render(payload)
    assert content.startswith(b"%PDF")
    assert len(content) > 1000


def test_context_text_carries_day_type_and_prev_year_line(monkeypatch):
    monkeypatch.setattr(daily_report.daily_context, "get_weather", lambda _d: None)
    # No dataset supplied -> the previous-year line degrades to the
    # "No historical data found" wording rather than being omitted.
    text = daily_report._context_text(date(2026, 9, 2), store="NM")
    assert "<b>Day type:</b> Mid-Week" in text
    assert "Previous year same day:</b> No historical data found" in text


def test_report_of_store_with_no_data_still_builds(db_session):
    payload = daily_report.build_daily_report_payload(db_session, "CHW", include_visuals=False)
    content = RENDERERS["xlsx"].render(payload)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "Date Wise Store Summary" in wb.sheetnames
