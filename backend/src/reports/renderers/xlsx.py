"""XLSX renderer -- a "Summary" sheet (title/timestamp/filters/KPI grid),
one sheet per TableBlock, and chart/gauge blocks embedded as images
(openpyxl already a project dependency, no new library needed here).
"""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font

from src.reports.chart_image import fig_dict_to_png, gauge_spec_to_png
from src.reports.models import (
    ChartBlock,
    DataQualityBlock,
    FiltersSummaryBlock,
    ForecastSummaryBlock,
    GaugeBlock,
    KpiGridBlock,
    ReportPayload,
    TableBlock,
    TextBlock,
    TitleBlock,
)

_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _sheet_name(title: str, used: set[str]) -> str:
    name = _INVALID_SHEET_CHARS.sub(" ", title)[:31].strip() or "Sheet"
    candidate = name
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = name[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _add_image(ws, png_bytes: bytes, anchor: str) -> None:
    img = XLImage(io.BytesIO(png_bytes))
    img.width, img.height = 480, 260
    ws.add_image(img, anchor)


def render(payload: ReportPayload) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    used_names = {"Summary"}

    bold = Font(bold=True)
    summary["A1"] = payload.meta.title
    summary["A1"].font = Font(bold=True, size=14)
    summary["A2"] = "Generated"
    summary["B2"] = payload.meta.generated_at
    summary["A3"] = "Filters"
    summary["B3"] = payload.meta.filters_summary_text
    row = 5

    for block in payload.blocks:
        if isinstance(block, TitleBlock):
            continue
        if isinstance(block, TextBlock):
            summary.cell(row=row, column=1, value=block.heading or "Notes").font = bold
            row += 1
            summary.cell(row=row, column=1, value=block.text)
            row += 2
        elif isinstance(block, FiltersSummaryBlock):
            summary.cell(row=row, column=1, value=block.heading).font = bold
            row += 1
            summary.cell(row=row, column=1, value=block.text)
            row += 2
        elif isinstance(block, KpiGridBlock):
            summary.cell(row=row, column=1, value=block.heading or "KPI Summary").font = bold
            row += 1
            for item in block.items:
                summary.cell(row=row, column=1, value=item.label)
                summary.cell(row=row, column=2, value=item.value)
                row += 1
            row += 1
        elif isinstance(block, ChartBlock):
            summary.cell(row=row, column=1, value=block.title).font = bold
            row += 1
            _add_image(summary, fig_dict_to_png(block.figure), f"A{row}")
            row += 16
        elif isinstance(block, GaugeBlock):
            title = str(block.spec.get("title", "Gauge"))
            summary.cell(row=row, column=1, value=title).font = bold
            row += 1
            _add_image(summary, gauge_spec_to_png(block.spec), f"A{row}")
            row += 16
        elif isinstance(block, TableBlock):
            ws = wb.create_sheet(_sheet_name(block.title, used_names))
            for col_idx, col_name in enumerate(block.columns, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = bold
            for row_idx, table_row in enumerate(block.rows, start=2):
                for col_idx, col_name in enumerate(block.columns, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=table_row.get(col_name))
        elif isinstance(block, DataQualityBlock):
            summary.cell(row=row, column=1, value=block.heading).font = bold
            row += 1
            for item in block.items:
                summary.cell(row=row, column=1, value=f"- {item}")
                row += 1
            row += 1
        elif isinstance(block, ForecastSummaryBlock):
            summary.cell(row=row, column=1, value=block.heading).font = bold
            row += 1
            for item in block.items:
                summary.cell(row=row, column=1, value=item.label)
                summary.cell(row=row, column=2, value=item.value)
                row += 1
            if block.note:
                summary.cell(row=row, column=1, value=block.note)
                row += 1
            row += 1

    summary.cell(row=row + 1, column=1, value=payload.meta.workbook_note)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
