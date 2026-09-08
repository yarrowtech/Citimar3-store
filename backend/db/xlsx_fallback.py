"""TEMPORARY local-testing fallback: when MONGODB_URI isn't configured (no
real MongoDB set up yet), db/session.py falls back to this module instead of
raising -- Daily Operations reads/writes TEST_DAILY_DASHBOARD.xlsx directly,
the same file/sheet layout the app used before the Postgres/MongoDB
migrations (see CLAUDE.md's Daily Dashboard section for that history). This
is scaffolding for testing without a database account, not a third
supported backend -- the moment MONGODB_URI is set to a real value,
db/session.py stops falling back here automatically, no code change needed
anywhere else.

Implements exactly the pymongo Collection surface src/daily_dashboard_store.py
actually calls (find/find_one/insert_one/find_one_and_update/update_one/
delete_one/aggregate's narrow $match+$group+$sum shape, create_index as a
no-op) -- not a general MongoDB emulator, so it only works because this
project controls both call sites.

`_id` is the worksheet row number (row 2 = the first data row) -- the
original xlsx-era design's row-identity convention, restored on purpose
here: delete_one blanks the row's STORE cell rather than physically removing
it (find/list functions skip blank-STORE rows as not-a-real-row), which is
what keeps `_id` stable across deletes. A real row removal would silently
renumber every later row's `_id` out from under any cached frontend state.
next_id (db/models.py) calling db["counters"] is special-cased below to hand
back "the next row number in <target sheet>" instead of a stored sequence,
so `_id` always equals the row a document actually gets written to -- no
separate counter to keep in sync with the sheet's real contents.
"""
from __future__ import annotations

import json
import threading
from datetime import date as date_cls
from typing import Any, Iterator

import openpyxl

from config.settings import DAILY_DASHBOARD_XLSX_PATH

_LOCK = threading.Lock()

_SHEETS = {
    "bills": "BILLS",
    "footfall": "FOOTFALL",
    "nob": "NOB",
    "targets": "TARGETS",
}

_COLUMNS: dict[str, list[tuple[str, str]]] = {
    "bills": [
        ("store_code", "STORE"), ("entry_date", "DATE"), ("bill_time", "BILL TIME STAMP"),
        ("net_amount", "NET AMOUNT"), ("bill_quantity", "BILL QUANTITY"), ("time_slot", "TIME SLOT"),
    ],
    "footfall": [
        ("store_code", "STORE"), ("entry_date", "DATE"), ("entry_time", "TIME"),
        ("footfall", "FOOTFALL"), ("time_slot", "TIME SLOT"),
    ],
    "nob": [
        ("store_code", "STORE"), ("entry_date", "DATE"), ("entry_time", "TIME"),
        ("nob", "NOB"), ("time_slot", "TIME SLOT"),
    ],
    "targets": [
        ("store_code", "STORE"), ("entry_date", "DATE"), ("sales_target", "SALES TARGET"),
        ("net_sales", "NET SALES"), ("remaining", "REMAINING"), ("footfall", "FOOTFALL"),
        ("nob", "NOB"), ("atv", "ATV"), ("rpv", "RPV"), ("basket_size", "BASKET SIZE"),
        ("conversion_pct", "CONVERSION %"), ("achievement_pct", "ACHIEVEMENT %"),
        ("reason", "REASON FOR TARGET VARIANCE"),
        # A manager's hand-entered KPI overrides (targets.overrides sub-doc in
        # Mongo) flattened to one JSON-string cell here -- the fallback has no
        # nested-document type, and this keeps set_kpi_override's dotted
        # `$set: {"overrides.atv": ...}` / `$unset` working locally.
        ("overrides", "KPI OVERRIDES (JSON)"),
    ],
}


def _cell_to_value(field: str, raw: Any) -> Any:
    if field == "entry_date":
        if raw is None:
            return None
        if hasattr(raw, "date"):
            raw = raw.date()
        return raw.isoformat()
    if field == "overrides":
        if raw in (None, ""):
            return {}
        try:
            return json.loads(raw)
        except (ValueError, TypeError):
            return {}
    return raw


def _value_to_cell(field: str, value: Any) -> Any:
    if field == "entry_date" and isinstance(value, str):
        return date_cls.fromisoformat(value)
    if field == "overrides":
        return json.dumps(value) if value else None
    return value


def _row_to_doc(collection: str, row_idx: int, values: tuple) -> dict | None:
    if values[0] in (None, ""):  # STORE blank -- soft-deleted or never-used scaffold row
        return None
    doc: dict = {"_id": row_idx}
    for (field, _header), raw in zip(_COLUMNS[collection], values):
        doc[field] = _cell_to_value(field, raw)
    return doc


def _set_cell(ws, row: int, col: int, value: Any) -> None:
    # openpyxl's ws.cell(row, col, value=None) is a no-op, not "clear this
    # cell" -- value=None there means "no value was passed," so clearing a
    # cell (or writing a genuinely-None field like an unset `reason`) needs
    # the two-step form: fetch the cell object, then assign .value directly.
    ws.cell(row=row, column=col).value = value


class _UpdateResult:
    def __init__(self, deleted_count: int = 0):
        self.deleted_count = deleted_count


class _Cursor:
    def __init__(self, docs: list[dict]):
        self._docs = docs

    def sort(self, field: str, direction: int = 1) -> "_Cursor":
        self._docs.sort(key=lambda d: (d.get(field) is None, d.get(field)), reverse=direction < 0)
        return self

    def __iter__(self):
        return iter(self._docs)


class XlsxCollection:
    def __init__(self, name: str):
        self.name = name

    def _open(self):
        wb = openpyxl.load_workbook(DAILY_DASHBOARD_XLSX_PATH)
        return wb, wb[_SHEETS[self.name]]

    def _iter_docs(self, ws) -> Iterator[dict]:
        fields = _COLUMNS[self.name]
        for row_idx in range(2, ws.max_row + 1):
            values = tuple(ws.cell(row=row_idx, column=c + 1).value for c in range(len(fields)))
            doc = _row_to_doc(self.name, row_idx, values)
            if doc is not None:
                yield doc

    @staticmethod
    def _matches(doc: dict, filter_: dict) -> bool:
        return all(doc.get(k) == v for k, v in filter_.items())

    def find_one(self, filter_: dict) -> dict | None:
        with _LOCK:
            _wb, ws = self._open()
            return next((d for d in self._iter_docs(ws) if self._matches(d, filter_)), None)

    def find(self, filter_: dict) -> _Cursor:
        with _LOCK:
            _wb, ws = self._open()
            return _Cursor([d for d in self._iter_docs(ws) if self._matches(d, filter_)])

    def insert_one(self, doc: dict) -> None:
        with _LOCK:
            wb, ws = self._open()
            row_idx = doc["_id"]
            for col_idx, (field, _header) in enumerate(_COLUMNS[self.name], start=1):
                _set_cell(ws, row_idx, col_idx, _value_to_cell(field, doc.get(field)))
            wb.save(DAILY_DASHBOARD_XLSX_PATH)

    def delete_one(self, filter_: dict) -> _UpdateResult:
        with _LOCK:
            wb, ws = self._open()
            for doc in self._iter_docs(ws):
                if self._matches(doc, filter_):
                    _set_cell(ws, doc["_id"], 1, None)
                    wb.save(DAILY_DASHBOARD_XLSX_PATH)
                    return _UpdateResult(deleted_count=1)
            return _UpdateResult(deleted_count=0)

    def _apply_update(self, ws, doc: dict, update: dict) -> dict:
        """Apply a narrow subset of Mongo update operators ($set, $unset) --
        including dotted `overrides.<field>` keys, which map onto the single
        flattened JSON overrides cell -- and return the post-update doc."""
        field_to_col = {f: i + 1 for i, (f, _h) in enumerate(_COLUMNS[self.name])}
        row_idx = doc["_id"]
        result = {**doc}
        overrides = dict(doc.get("overrides") or {})
        overrides_touched = False

        for field, value in update.get("$set", {}).items():
            if field.startswith("overrides."):
                overrides[field.split(".", 1)[1]] = value
                overrides_touched = True
            elif field in field_to_col:
                _set_cell(ws, row_idx, field_to_col[field], _value_to_cell(field, value))
                result[field] = value
        for field in update.get("$unset", {}):
            if field.startswith("overrides."):
                overrides.pop(field.split(".", 1)[1], None)
                overrides_touched = True

        if overrides_touched and "overrides" in field_to_col:
            _set_cell(ws, row_idx, field_to_col["overrides"], _value_to_cell("overrides", overrides))
            result["overrides"] = overrides
        return result

    def find_one_and_update(self, filter_: dict, update: dict, return_document=None) -> dict | None:
        with _LOCK:
            wb, ws = self._open()
            for doc in self._iter_docs(ws):
                if self._matches(doc, filter_):
                    result = self._apply_update(ws, doc, update)
                    wb.save(DAILY_DASHBOARD_XLSX_PATH)
                    return result
            return None

    def update_one(self, filter_: dict, update: dict) -> None:
        with _LOCK:
            wb, ws = self._open()
            for doc in self._iter_docs(ws):
                if self._matches(doc, filter_):
                    self._apply_update(ws, doc, update)
                    wb.save(DAILY_DASHBOARD_XLSX_PATH)
                    return

    def aggregate(self, pipeline: list[dict]) -> list[dict]:
        match = next((stage["$match"] for stage in pipeline if "$match" in stage), {})
        group = next((stage["$group"] for stage in pipeline if "$group" in stage), None)
        with _LOCK:
            _wb, ws = self._open()
            docs = [d for d in self._iter_docs(ws) if self._matches(d, match)]
        if group is None or not docs:
            return []
        result: dict = {}
        for key, spec in group.items():
            if key == "_id":
                continue
            field = spec["$sum"].lstrip("$")
            result[key] = sum((d.get(field) or 0) for d in docs)
        return [result]

    def create_index(self, *args, **kwargs) -> None:
        pass  # xlsx has no real indexes -- no-op, matches Mongo's idempotent create_index


class _CountersCollection:
    """See module docstring -- hands back "next row number in <target
    sheet>" rather than a stored sequence, so ids and actual sheet contents
    can never drift apart."""

    def find_one_and_update(self, filter_: dict, update: dict, upsert: bool = False, return_document=None) -> dict:
        target = filter_["_id"]
        with _LOCK:
            wb = openpyxl.load_workbook(DAILY_DASHBOARD_XLSX_PATH)
            ws = wb[_SHEETS[target]]
            return {"_id": target, "seq": ws.max_row + 1}


class XlsxDatabase:
    """Duck-typed pymongo.database.Database replacement -- see module
    docstring for what it does and does not implement."""

    def __getitem__(self, name: str):
        if name == "counters":
            return _CountersCollection()
        return XlsxCollection(name)
