"""DOCX renderer via python-docx: headings/paragraphs, add_table for
KPI/table blocks, add_picture for rasterized chart/gauge images.
"""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Inches, Pt

from config.settings import PROJECT_ROOT
from src.reports.chart_image import fig_dict_to_png, gauge_spec_to_png
from src.reports.models import (
    ChartBlock,
    DataQualityBlock,
    FiltersSummaryBlock,
    ForecastSummaryBlock,
    GaugeBlock,
    KpiGridBlock,
    ReportPayload,
    TableBlock,
    TextBlock,
    TitleBlock,
)

_LOGO_PATH = PROJECT_ROOT / "img" / "logo" / "CitiMart_logo_2.png"


def _add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value


def _add_data_table(doc: Document, columns: list[str], rows: list[dict]) -> None:
    table = doc.add_table(rows=1, cols=max(len(columns), 1))
    table.style = "Light Grid Accent 1"
    for idx, col in enumerate(columns):
        table.rows[0].cells[idx].text = str(col)
    for row in rows:
        cells = table.add_row().cells
        for idx, col in enumerate(columns):
            cells[idx].text = "" if row.get(col) is None else str(row.get(col))


def render(payload: ReportPayload) -> bytes:
    doc = Document()

    if _LOGO_PATH.exists():
        doc.add_picture(str(_LOGO_PATH), width=Inches(1.8))

    title = doc.add_heading(payload.meta.title, level=0)
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Generated: {payload.meta.generated_at}").italic = True
    doc.add_paragraph(f"Filters: {payload.meta.filters_summary_text}")
    doc.add_paragraph()

    for block in payload.blocks:
        if isinstance(block, TitleBlock):
            continue
        if isinstance(block, TextBlock):
            if block.heading:
                doc.add_heading(block.heading, level=1)
            doc.add_paragraph(block.text)
        elif isinstance(block, FiltersSummaryBlock):
            doc.add_heading(block.heading, level=1)
            doc.add_paragraph(block.text)
        elif isinstance(block, KpiGridBlock):
            if block.heading:
                doc.add_heading(block.heading, level=1)
            _add_kv_table(doc, [(i.label, i.value) for i in block.items])
        elif isinstance(block, ChartBlock):
            doc.add_heading(block.title, level=1)
            doc.add_picture(io.BytesIO(fig_dict_to_png(block.figure)), width=Inches(6))
        elif isinstance(block, GaugeBlock):
            title_text = str(block.spec.get("title", "Gauge"))
            doc.add_heading(title_text, level=1)
            doc.add_picture(io.BytesIO(gauge_spec_to_png(block.spec)), width=Inches(4))
        elif isinstance(block, TableBlock):
            doc.add_heading(block.title, level=1)
            _add_data_table(doc, block.columns, block.rows)
        elif isinstance(block, DataQualityBlock):
            doc.add_heading(block.heading, level=1)
            for item in block.items:
                doc.add_paragraph(item, style="List Bullet")
        elif isinstance(block, ForecastSummaryBlock):
            doc.add_heading(block.heading, level=1)
            _add_kv_table(doc, [(i.label, i.value) for i in block.items])
            if block.note:
                doc.add_paragraph(block.note).italic = True
        doc.add_paragraph()

    footer_p = doc.add_paragraph()
    footer_run = footer_p.add_run(payload.meta.workbook_note)
    footer_run.italic = True
    footer_run.font.size = Pt(9)

    section = doc.sections[0]
    section.footer.paragraphs[0].text = payload.meta.workbook_note

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
